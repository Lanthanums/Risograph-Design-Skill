"""Deterministic poster fixtures and negative geometry/policy regression tests."""
import argparse
import copy
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

ROOT=Path(__file__).resolve().parents[1]
sys.dont_write_bytecode=True
sys.path.insert(0,str(ROOT/'scripts'))
from pptx import Presentation
from pptx.util import Pt
from PIL import Image
import build_poster as bp
from canvas import DECK_CANVAS, poster_canvas, POSTER_PRESETS
import compose
import layout_check as lc
import qa
import riso_kit as rk


def part_hash(path):
    with zipfile.ZipFile(path) as z:
        return hashlib.sha256(b''.join(n.encode()+z.read(n) for n in sorted(z.namelist()))).hexdigest()


def geometry(slide,canvas,style,media):
    rep=qa.Report()
    args=argparse.Namespace(edge_tol=1.5,text_protection=9,media=str(media),canvas=canvas,style=style)
    qa.geometry_checks(rep,slide,1,args)
    return [r for r in rep.items if not r['ok'] and r['severity']=='fail']


class PosterTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(prefix='riso_poster_')
        self.work=Path(self.tmp.name)
        self.env=patch.dict(os.environ,{'RISO_DESIGN_STATE_DIR':str(self.work/'state')})
        self.env.start()

    def tearDown(self):
        self.env.stop()
        self.tmp.cleanup()

    def build(self,spec,name='poster'):
        out=self.work/name/'poster.pptx'
        bp.build(spec,str(out),self.work)
        style=json.loads((out.parent/'style.json').read_text(encoding='utf-8'))
        return out,Presentation(out),style

    def check_artifact(self,out,prs,style):
        canvas=poster_canvas({'width':prs.slide_width.pt,'height':prs.slide_height.pt})
        self.assertEqual(len(prs.slides),1)
        self.assertEqual(geometry(prs.slides[0],canvas,style,out.parent/'media'),[])
        report=lc.check_slide(prs.slides[0],1,canvas=canvas,
                    primary=style['dialect']['primary'],secondary=style['dialect']['secondary'],
                    poster_policy=style['poster'])
        self.assertTrue(report['pass'],report['fails'])
        for tool in ('qa.py','layout_check.py'):
            cmd=[sys.executable,str(ROOT/'scripts'/tool),'--deck',str(out),
                 '--json',str(out.parent/(tool+'.json'))]
            if tool=='qa.py':cmd+=['--media',str(out.parent/'media')]
            p=subprocess.run(cmd,capture_output=True,text=True,encoding='utf-8',errors='replace')
            self.assertEqual(p.returncode,0,p.stdout+p.stderr)

    def test_fixtures_and_determinism(self):
        for f in sorted((ROOT/'tests'/'posters').glob('*.json')):
            with self.subTest(f=f.name):
                spec=json.loads(f.read_text(encoding='utf-8'))
                out,prs,style=self.build(spec,f.stem)
                expected=poster_canvas(spec['canvas'])
                self.assertAlmostEqual(prs.slide_width.pt,expected.width,places=3)
                self.assertAlmostEqual(prs.slide_height.pt,expected.height,places=3)
                self.check_artifact(out,prs,style)
                out2,_,_=self.build(spec,f.stem+'-again')
                self.assertEqual(part_hash(out),part_hash(out2))
                self.assertEqual(out.read_bytes(),out2.read_bytes())
                self.assertEqual((out.parent/'media'/'bg1.jpg').read_bytes(),
                                 (out2.parent/'media'/'bg1.jpg').read_bytes())

    def test_family_priority_and_custom_canvas(self):
        self.assertEqual(bp.choose_family({'image':'x'},'type_led',1),('type_led','user'))
        self.assertEqual(bp.choose_family({'date':'x'},None,1),('editorial_event','semantics'))
        self.assertEqual(bp.choose_family({'title':'x'},None,8),bp.choose_family({'title':'x'},None,8))
        c=poster_canvas({'width':200,'height':300,'unit':'mm'})
        self.assertAlmostEqual(c.aspect_ratio,2/3)
        self.assertEqual(poster_canvas({'width':600,'ratio':'2:3'}).height,900)
        for bad in ({'width':0},{'height':float('nan')},{'ratio':0},{'unit':'px'},
                    {'width':100,'height':100,'ratio':'4:5'}):
            with self.assertRaises((ValueError,ZeroDivisionError)):
                poster_canvas(bad)
        spec={'canvas':{'width':600,'ratio':'2:3'},'family':'type_led',
              'meta':{'elements':False,'dialect':['C','A']},'content':{'title':'CUSTOM'}}
        out,prs,style=self.build(spec)
        self.check_artifact(out,prs,style)
        spec['content'].update({'blocks':['Supplied supporting block'],'text_bleed':['block']})
        out,prs,style=self.build(spec,'block-trim')
        self.check_artifact(out,prs,style)
        self.assertTrue(any(r['role']=='block' and r['rect'][0]==0
                            for r in style['poster']['roles']))
        spec['content']['text_bleed']=['image']
        with self.assertRaises(ValueError):self.build(spec,'invalid-bleed-role')

    def test_chrome_missing_font_and_overflow_fail(self):
        spec=json.loads((ROOT/'tests'/'posters'/'data_en_square.json').read_text())
        bad=copy.deepcopy(spec);bad['meta']['chrome']=['number']
        with self.assertRaises(ValueError):self.build(bad)
        bad=copy.deepcopy(spec);bad['meta']['fonts']={'heading':'NoSuchRisoFont'}
        with self.assertRaises(ValueError):self.build(bad)
        bad=copy.deepcopy(spec);bad['content']['body']='Too much source copy. '*500
        with self.assertRaises(ValueError):self.build(bad)
        out,prs,style=self.build(spec)
        canvas=poster_canvas(spec['canvas'])
        rk.text(prs.slides[0],500,710,400,50,[dict(runs=[('OFF CANVAS',{})],
                font='Georgia',size=30,color='#2B2B2B')])
        self.assertTrue(geometry(prs.slides[0],canvas,style,out.parent/'media'))
        rk.text(prs.slides[0],50,550,100,24,[dict(runs=[('01 / 01',{})],
                font='Georgia',size=18,color='#2B2B2B')])
        report=lc.check_slide(prs.slides[0],1,canvas=canvas,poster_policy=style['poster'])
        self.assertTrue(any('chrome' in s for s in report['fails']),report['fails'])

    def test_quote_image_treatment_and_full_bleed(self):
        img=self.work/'user.png'
        Image.new('RGB',(220,300),(30,80,120)).save(img)
        for family in ('image_led','manifesto_quote'):
            spec={'canvas':'a3_portrait','family':family,'meta':{'elements':False,'dialect':['C','A']},
                  'content':{'title':'LOOK AGAIN'}}
            if family=='image_led':spec['content']['image']={'path':str(img)}
            else:spec['content']['quote']={'text':'A question is a place to begin.','attribution':'Supplied fixture quote'}
            out,prs,style=self.build(spec,family)
            self.check_artifact(out,prs,style)
            if family=='image_led':
                out2,_,_=self.build(spec,family+'-again')
                self.assertEqual(part_hash(out),part_hash(out2))
                pic=next(s for s in prs.slides[0].shapes if s.name.startswith('user-image['))
                pic._element.getparent().remove(pic._element)
                replacement=prs.slides[0].shapes.add_picture(str(img),Pt(36),Pt(500),Pt(300),Pt(300))
                replacement.name='user-image[duotone]'
                report=lc.check_slide(prs.slides[0],1,canvas=poster_canvas(),poster_policy=style['poster'])
                self.assertTrue(any('provenance' in s for s in report['fails']),report['fails'])
                replacement.name='Picture 99'  # renaming must not bypass treatment checks
                report=lc.check_slide(prs.slides[0],1,canvas=poster_canvas(),poster_policy=style['poster'])
                self.assertTrue(any('provenance' in s for s in report['fails']),report['fails'])
        spec={'family':'image_led','meta':{'elements':False,'dialect':['C','A']},
              'content':{'title':'FULL BLEED','image':{'path':str(img)},
                         'full_bleed_reason':'The supplied image is the content protagonist'}}
        out,prs,style=self.build(spec,'full')
        self.check_artifact(out,prs,style)
        style['poster']['full_bleed_reason']=None
        report=lc.check_slide(prs.slides[0],1,canvas=poster_canvas(),poster_policy=style['poster'])
        self.assertFalse(report['pass'])

    def test_builder_succeeds_without_renderer(self):
        spec=self.work/'spec.json'
        spec.write_text(json.dumps({'content':{'title':'EDITABLE'},'meta':{'elements':False}}))
        cmd=[sys.executable,str(ROOT/'scripts'/'build_poster.py'),'--spec',str(spec),
             '--out',str(self.work/'cli'/'poster.pptx')]
        p=subprocess.run(cmd,capture_output=True,text=True,encoding='utf-8',errors='replace')
        self.assertEqual(p.returncode,0,p.stdout+p.stderr)
        self.assertTrue((self.work/'cli'/'poster.pptx').is_file())


if __name__=='__main__':unittest.main()
