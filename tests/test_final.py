"""Final behavior: graphic composition, opt-outs, diagnostics and cheap queries."""
import sys
sys.dont_write_bytecode=True
import argparse
import copy
import json
import os
from pathlib import Path
import subprocess
import tempfile
import unittest
from unittest.mock import patch

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
from pptx import Presentation
import artifact_io as aio
import build_poster as bp
import build_deck as bd
import compose
import layout as lay
import layout_check as lc
import qa
import riso_kit as rk
import textmetrics as tm
from canvas import poster_canvas


class FinalTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.work=Path(self.temp.name)
        self.env=patch.dict(os.environ,{'RISO_DESIGN_STATE_DIR':str(self.work/'state')})
        self.env.start()

    def tearDown(self):
        self.env.stop();self.temp.cleanup()

    def spec(self,dialect='C'):
        return {'artifact':'poster','meta':{'lang':'en','palette':'C','seed':17,'dialect':[dialect,'C']},
                'layout':{'graphic':'band'},'content':{'title':'SHAPE IDEAS','subtitle':'Work with a question.',
                'body':'Sketch the parts. Explore their relationships. Test what changes.'}}

    def build(self,spec,name):
        out=self.work/(name+'.pptx');bp.build(spec,out)
        return out,Presentation(out),aio.read_style(out)

    def test_composition_across_dialects_and_determinism(self):
        for dialect in 'ABCDEF':
            with self.subTest(dialect=dialect):
                spec=self.spec(dialect);out,prs,style=self.build(spec,dialect)
                group=[e for e in style['elements'] if e.get('group')=='composition']
                self.assertTrue(style['poster']['graphic']['placed'])
                self.assertTrue(group)
                self.assertTrue(all('error' not in e for e in style['elements']))
                canvas=poster_canvas(spec.get('canvas'))
                self.assertGreater(max(lay.area(e['rect']) for e in group),canvas.area*.015)
                result=lc.check_slide(prs.slides[0],1,canvas=canvas,primary=dialect,secondary='C',poster_policy=style['poster'])
                self.assertFalse(result['fails'],result['fails'])
                report=qa.Report()
                qa.geometry_checks(report,prs.slides[0],1,argparse.Namespace(edge_tol=1.5,text_protection=9,
                     media=str(aio.paths(out)[2]),canvas=canvas,style=style))
                self.assertFalse([r for r in report.items if not r['ok'] and r['severity']=='fail'])
                if dialect=='C':
                    again,_,_=self.build(spec,'again')
                    self.assertEqual(out.read_bytes(),again.read_bytes())

    def test_pure_type_and_empty_pool_do_not_reserve_band(self):
        for setting in ({'elements':False},{'element_pool':[]}):
            spec=self.spec();spec.pop('layout');spec['meta'].update(setting)
            plan=bp.plan(spec)['plan']
            self.assertIsNone(plan['graphic_rect'])
            _,_,style=self.build(spec,'pure')
            self.assertEqual(style['elements'],[])

    def test_explicit_graphic_filters_and_impossible_band(self):
        spec=self.spec();spec['meta']['graphic_element']='arch'
        _,_,style=self.build(spec,'chosen')
        self.assertEqual(next(e['element'] for e in style['elements'] if e.get('group')=='composition'),'arch')
        for mutation in ({'exclude_elements':['arch']},{'elements':False},{'element_pool':[]}):
            bad=copy.deepcopy(spec);bad['meta'].update(mutation)
            with self.assertRaises(ValueError):self.build(bad,'invalid')
        bad=self.spec();bad['meta']['element_pool']=['stripe_band']
        with self.assertRaisesRegex(ValueError,'graphic band'):
            self.build(bad,'no-main-shape')

    def test_multiline_display_is_protected_from_accents(self):
        spec=json.loads((ROOT/'tests/posters/type_en_a3.json').read_text(encoding='utf-8'))
        _,prs,style=self.build(spec,'multiline')
        title=next(r for r in style['poster']['roles'] if r['role']=='title')
        shapes=[s for s in prs.slides[0].shapes if s.name in title['shapes']]
        bounds=(min(s.left.pt for s in shapes),min(s.top.pt for s in shapes),
                max(tm.ink_rect(s)[2] for s in shapes),max(s.top.pt+s.height.pt for s in shapes))
        for e in style['elements']:self.assertFalse(lay.intersection(e['rect'],bounds),e)

    def test_display_offsets_between_previous_extremes(self):
        for size in (24,34,70,100,148,300):
            normal=rk.misreg_offsets(size,'print')[0]
            self.assertGreater(normal,min(size*.025,3.5))
            self.assertLess(normal,rk.misreg_offsets(size)[0])
            self.assertEqual(rk.misreg_offsets(size,'reverse')[0],min(size*.075,12))

    def test_summary_purity_and_curated_catalogue(self):
        spec=self.work/'poster.json';spec.write_text(json.dumps(self.spec()),encoding='utf-8')
        before=set(self.work.rglob('*'))
        result=subprocess.run([sys.executable,str(ROOT/'scripts/build_poster.py'),'--spec',str(spec),'--plan','--summary'],
                              capture_output=True,text=True,encoding='utf-8')
        self.assertEqual(result.returncode,0,result.stderr)
        self.assertNotIn('SHAPE IDEAS',result.stdout)
        self.assertIsNotNone(json.loads(result.stdout)['graphic_rect'])
        self.assertEqual(before,set(self.work.rglob('*')))
        result=subprocess.run([sys.executable,str(ROOT/'scripts/elements.py'),'catalogue','--dialect','D','--role','hero','--safe','--json'],
                              capture_output=True,text=True,encoding='utf-8')
        self.assertEqual(result.returncode,0,result.stderr)
        names={e['name'] for e in json.loads(result.stdout)}
        self.assertEqual(names,{n for role,n in compose.element_pool('D','D') if role=='hero'})
        self.assertFalse(names & (compose.TEXT_BEARING | compose.NUMBER_BEARING))

    def test_metric_cache_matches_uncached_and_stays_bounded(self):
        theme=rk.Theme(lang='en')
        for text in ('Spacing matters','123.45',''):
            self.assertEqual(tm.advance(text,theme.f_body,23),tm.advance.__wrapped__(text,theme.f_body,23))
        self.assertIsNotNone(tm.advance.cache_info().maxsize)

    def test_timeline_graphics_are_not_decorative_landing_zones(self):
        source=json.loads((ROOT/'tests/fixtures/course_zh_en.json').read_text(encoding='utf-8'))
        page=next(p for p in source['pages'] if p['archetype']=='timeline')
        spec={'meta':{'lang':'zh','seed':source['meta']['seed'],'dialect':['B','A'],'elements':True},'pages':[page]}
        out=self.work/'timeline.pptx';bd.build(spec,out)
        slide=Presentation(out).slides[0]
        lines=[lay.inflate(lay.obb(s),2) for s in slide.shapes if lay.line_segment(s)]
        self.assertTrue(lines)
        for s in slide.shapes:
            if s.name.startswith('el['):
                self.assertFalse(any(lay.intersection(lay.obb(s),r) for r in lines),s.name)


if __name__=='__main__':unittest.main()
