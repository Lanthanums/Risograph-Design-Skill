"""Poster composition families. Content roles, measured type, no deck archetypes."""
import sys
sys.dont_write_bytecode = True
import random
import re

import riso_kit as rk
import textmetrics as tm
import typefit as tf

FAMILIES = ('type_led','image_led','editorial_event','data_poster','manifesto_quote')
TEXT_ROLES = ('title','subtitle','body','block','metadata','cta','key_number',
              'key_number_label','quote','attribution','footer','block_title','comparison','process')
TITLE_MIN_SHARE = 0.05
TITLE_BODY_RATIO = 1.8


def _present(value):
    return value is not None and str(value).strip()!=''


def choose_family(content, requested, seed):
    if requested:
        if requested not in FAMILIES:
            raise ValueError(f'Unknown poster family {requested!r}; choose {FAMILIES}')
        return requested,'user'
    if content.get('image'):
        return 'image_led','semantics'
    if _present(content.get('key_number')) or content.get('data'):
        return 'data_poster','semantics'
    if content.get('date') or content.get('venue'):
        return 'editorial_event','semantics'
    if content.get('quote'):
        return 'manifesto_quote','semantics'
    return random.Random(seed).choice(('type_led','editorial_event')),'seed'



lines_of = tf.lines_of  # old public helper stays available

LAYOUT_DEFAULTS={'columns':'auto','structure':'rules','alignment':'l',
                 'footer_anchor':'content','vertical_alignment':'balanced','graphic':'auto',
                 'graphic_position':'auto'}


def layout_options(raw=None):
    raw=raw or {}
    unknown=set(raw)-set(LAYOUT_DEFAULTS)-{'title_size','body_size'}
    if unknown:raise ValueError(f'Unknown poster layout fields: {sorted(unknown)}')
    out=dict(LAYOUT_DEFAULTS,**raw)
    allowed={'columns':('auto',1,2),'structure':('rules','none'),'alignment':('l','c','r'),
             'footer_anchor':('content','edge'),'vertical_alignment':('balanced','top'),
             'graphic':('auto','band','none'),
             'graphic_position':('auto','after_title','after_content')}
    for key,values in allowed.items():
        if out[key] not in values or isinstance(out[key],bool):
            raise ValueError(f'poster layout.{key} must be one of {values}')
    import math
    for key in ('title_size','body_size'):
        if key in out and (isinstance(out[key],bool) or not isinstance(out[key],(int,float)) or not math.isfinite(out[key]) or out[key]<=0):
            raise ValueError(f'poster layout.{key} must be positive and finite')
    return out


class Typesetter:
    def __init__(self,slide,theme,canvas,text_bleed=(),alignment='l'):
        self.slide,self.th,self.canvas=slide,theme,canvas
        self.text_bleed=set(text_bleed)
        self.alignment=alignment
        self.roles=[]

    def measure(self,role,value,w,max_h,size,display=False,minimum=16,bold=False,index=None):
        family=(self.th.f_heading if re.search(r'[\u3400-\u9fff]',str(value))
                else self.th.f_display) if display else self.th.f_body
        if role in self.text_bleed:w=self.canvas.width
        offset=(lambda sz:max(rk.misreg_offsets(sz,'reverse' if self.th.on_red else 'print'))) if display else None
        return tf.fit(value,family,size,w,max_h,minimum=minimum,bold=bold,
                      role='poster '+role,block_index=index,offset=offset)

    def put_fit(self,role,f,x,y,w,display=False,index=None):
        if role in self.text_bleed:x,w=0,self.canvas.width
        names=[]
        if self.slide is not None:
            shapes=tf.draw(self.slide,x,y,w,f,self.th,display=display,
                           alignment=self.alignment,role='poster',block_index=index)
            for i,sh in enumerate(shapes):
                sh.name=f'poster[{role}]/block={index}/shape={i}/'+sh.name
                names.append(sh.name)
        self.roles.append({'role':role,'block_index':index,'order':len(self.roles),
                           'text':f.text,'font':f.font,'size':f.size,'lines':list(f.lines),
                           'rect':[x,y,x+w,y+f.height],'alignment':self.alignment,
                           'shapes':names})
        return y+f.height

    def put(self,role,value,x,y,w,max_h,size,display=False,minimum=16,bold=False,index=None):
        if value is None or str(value).strip()=='':return y
        return self.put_fit(role,self.measure(role,value,w,max_h,size,display,minimum,bold,index),
                            x,y,w,display,index)


def body_blocks(content):
    blocks=[]
    if content.get('body'):
        blocks.append(('body',str(content['body'])))
    for b in content.get('blocks',[]):
        if isinstance(b,str):
            blocks.append(('block',b))
        else:
            blocks.append(('block','\n'.join(str(b[k]) for k in ('title','text') if b.get(k))))
    for b in content.get('data',[]):
        blocks.append(('block','\n'.join(str(b[k]) for k in ('label','value') if k in b)))
    return blocks


def metadata(content):
    out=[str(content[k]) for k in ('date','venue') if content.get(k)]
    raw=content.get('metadata')
    if isinstance(raw,dict):
        out.extend(str(k)+': '+str(v) for k,v in raw.items())
    elif isinstance(raw,list):
        out.extend(str(v) for v in raw)
    elif raw:
        out.append(str(raw))
    return '\n'.join(out)



def compose_poster(slide,content,theme,canvas,family,seed,image_path=None,layout=None):
    """Plan measured regions first, print treatment next, decoration last.

    The same function plans without a slide. Unequal blocks consume only their
    measured height; balancing translates the composition by a bounded amount.
    No semantic relation or copy is synthesized.
    """
    options=layout_options(layout)
    x0,y0,x1,y1=canvas.safe
    w,h=x1-x0,y1-y0;short=min(canvas.width,canvas.height)
    gap=short*.025
    body_size=options.get('body_size',max(18,short*(.033 if family in ('type_led','editorial_event') else .026)))
    title_size=options.get('title_size',short*(.11 if family in ('data_poster','manifesto_quote') else .175))
    minimum=short*TITLE_MIN_SHARE
    setter=Typesetter(slide,theme,canvas,content.get('text_bleed',[]),options['alignment'])
    quote=content.get('quote')
    quote_text=quote.get('text') if isinstance(quote,dict) else quote
    attribution=quote.get('attribution') if isinstance(quote,dict) else content.get('attribution')
    number=content.get('key_number')
    number_text=number.get('value') if isinstance(number,dict) else number
    number_label=number.get('label') if isinstance(number,dict) else None
    if family=='image_led' and not image_path:raise ValueError('image_led requires a user image')
    if family=='data_poster' and not (_present(number_text) or content.get('data')):
        raise ValueError('data_poster requires key_number or data')
    foot=[]
    for role in ('cta','footer'):
        if content.get(role):
            f=setter.measure(role,content[role],w,h*.15,max(16,min(body_size,24)),minimum=16,bold=role=='cta')
            foot.append((role,f))
    footer_h=sum(f.height+gap for _,f in foot)
    bottom=y1-footer_h
    # Regions are measured into a schedule before any native object is drawn.
    regions=[];rules=[];cursor=y0
    def add(role,value,x,width,max_h,size,display=False,minimum=16,bold=False,index=None):
        nonlocal cursor
        if value is None or str(value).strip()=='':return
        f=setter.measure(role,value,width,max_h,size,display,minimum,bold,index)
        regions.append((role,f,x,cursor,width,display,index))
        cursor+=f.height+gap
    add('title',content['title'],x0,w,min(h*.38,bottom-cursor),title_size,True,minimum,True)
    if content.get('subtitle'):
        add('subtitle',content['subtitle'],x0,w,min(h*.17,bottom-cursor),body_size*1.2,minimum=17)
    graphic_at=cursor
    graphic_index=len(regions)
    if family=='data_poster' and _present(number_text):
        add('key_number',str(number_text),x0,w,min(h*.28,bottom-cursor),short*.24,True,minimum*1.5,True)
        if _present(number_label):add('key_number_label',number_label,x0,w,h*.12,body_size)
    if family=='manifesto_quote' and quote_text:
        add('quote',quote_text,x0,w,min(h*.36,bottom-cursor),short*.12,True,minimum,True)
        if attribution:add('attribution',attribution,x0,w,h*.10,body_size)
    meta=metadata(content)
    if meta:add('metadata',meta,x0,w,min(h*.20,bottom-cursor),body_size*1.05,minimum=17)
    modules=[];demotions=[]
    if _present(number_text) and family!='data_poster':
        modules.append([('key_number',str(number_text),None,True)])
        if _present(number_label):modules[-1].append(('key_number_label',str(number_label),None,False))
        demotions.append({'role':'key_number','to':'support','reason':f'{family} leads with another supplied role'})
    if quote_text and family!='manifesto_quote':
        modules.append([('quote',quote_text,None,False)])
        if attribution:modules[-1].append(('attribution',attribution,None,False))
        demotions.append({'role':'quote','to':'support','reason':f'{family} leads with another supplied role'})
    if content.get('body'):modules.append([('body',tf.source_text(content['body']),None,False)])
    for i,b in enumerate(content.get('blocks',[])):
        module=[]
        if isinstance(b,str):module=[('block',b,i,False)]
        else:
            if b.get('title'):module.append(('block_title',str(b['title']),i,True))
            if b.get('text'):module.append(('block',tf.source_text(b['text']),i,False))
        if module:modules.append(module)
    for i,b in enumerate(content.get('data',[])):
        modules.append([('block','\n'.join(str(b[k]) for k in ('label','value') if k in b),i,False)])
    # Diagrams require explicitly supplied labels and relations. An ordered
    # process list explicitly supplies order; comparison list supplies parallel
    # alternatives. No automatic flow is derived from ordinary prose.
    columns=options['columns']
    if columns=='auto':columns=2 if len(modules)>=3 and w>=600 and family!='image_led' else 1
    if family=='image_led':columns=1
    body_x=x0+w*.65 if family=='image_led' else x0
    body_w=w*.35 if family=='image_led' else w
    column_w=(body_w-gap*(columns-1))/columns
    cursors=[cursor]*columns
    # Contiguous column groups give a truthful column-major reading order.
    split=(len(modules)+1)//2 if columns==2 else len(modules)
    max_support=0
    for mi,module in enumerate(modules):
        col=1 if columns==2 and mi>=split else 0
        xx=body_x+col*(column_w+gap)
        for role,value,index,heading in module:
            yy=cursors[col]
            max_h=bottom-yy-gap
            size=body_size*1.2 if heading else body_size
            floor=17 if heading else 16
            f=setter.measure(role,value,column_w,max_h,size,False,floor,heading,index)
            regions.append((role,f,xx,yy,column_w,False,index))
            cursors[col]+=f.height+gap*.6
            max_support=max(max_support,f.size)
        if options['structure']=='rules' and mi<len(modules)-1:
            rules.append((xx,cursors[col]+gap*.3,column_w))
        cursors[col]+=gap
    if max_support:
        title_fit=regions[0][1]
        if title_fit.size<max_support*TITLE_BODY_RATIO:
            raise ValueError(f'poster title hierarchy cannot fit: title {title_fit.size:g}pt, support {max_support:g}pt; reallocate content/layout')
    end=max(cursors,default=cursor)
    diagrams=[]
    for kind in ('comparison','process'):
        labels=content.get(kind,[])
        if not labels:continue
        if len(labels)<2:raise ValueError(f'{kind} requires at least two supplied labels')
        count=len(labels)
        # Parallel framed alternatives express a supplied comparison. A vertical
        # connector expresses only the explicitly supplied process-list order.
        node_w=(body_w-gap*(count-1))/count if kind=='comparison' else body_w
        pad=gap*.55
        fits=[]
        for i,item in enumerate(labels):
            value='\n'.join(str(item[k]) for k in ('label','text') if item.get(k)) if isinstance(item,dict) else str(item)
            if not value:raise ValueError(f'{kind} block {i}: missing supplied label')
            available=(bottom-end-gap)/count-gap-2*pad if kind=='process' else bottom-end-gap-2*pad
            fits.append(setter.measure(kind,value,node_w-2*pad,available,body_size,minimum=16,index=i))
        heights=[f.height+2*pad for f in fits]
        row_h=max(heights)
        nodes=[]
        for i,f in enumerate(fits):
            xx=body_x+i*(node_w+gap) if kind=='comparison' else body_x
            yy=end if kind=='comparison' else end+sum(heights[:i])+gap*i
            hh=row_h if kind=='comparison' else heights[i]
            nodes.append([xx,yy,xx+node_w,yy+hh])
            regions.append((kind,f,xx+pad,yy+pad,node_w-2*pad,False,i))
        diagrams.append({'kind':kind,'nodes':nodes,'source_indices':list(range(count)),
                         'edges':[[i,i+1] for i in range(count-1)] if kind=='process' else []})
        end=max(n[3] for n in nodes)+gap
    image_rect=None
    if image_path:
        if content.get('full_bleed_reason'):image_rect=canvas.bounds
        elif family=='image_led':
            image_rect=(x0,cursor,x0+w*.60,bottom-gap)
            end=max(end,image_rect[3])
        else:
            image_rect=(x0,end,x1,min(bottom,end+h*.30))
            end=image_rect[3]+gap
        if image_rect[2]<=image_rect[0] or image_rect[3]<=image_rect[1]:
            raise ValueError('No remaining content room for supplied image')
    if end>bottom+1e-6:raise ValueError(f'poster content needs {end-y0:.1f}pt in {bottom-y0:.1f}pt; reallocate modules/columns')
    # Reserve a composition band from measured surplus BEFORE drawing. This is
    # a deliberate graphic/text relationship, not a hunt for leftover holes.
    # Existing images/diagrams carry the visual role in automatic mode.
    graphic_rect=None
    graphic_reason='disabled' if options['graphic']=='none' else 'content visual already present'
    eligible=not image_path and not diagrams
    if options['graphic']=='band' and not eligible:
        raise ValueError('graphic band conflicts with the supplied image/diagram; use auto or none')
    if eligible and options['graphic']!='none':
        room=bottom-end
        band_h=min(h*.22,short*.30)
        if room>=band_h+2*gap:
            delta=band_h+2*gap
            position=options['graphic_position']
            if position=='auto':
                position='after_content' if len(regions[0][1].lines)>1 and not foot else 'after_title'
            if position=='after_content':
                graphic_at=end
                graphic_index=len(regions)
            graphic_rect=(x0,graphic_at+gap,x1,graphic_at+gap+band_h)
            regions=[r if i<graphic_index else (*r[:3],r[3]+delta,*r[4:])
                     for i,r in enumerate(regions)]
            if position=='after_title':rules=[(x,y+delta,width) for x,y,width in rules]
            end+=delta
            graphic_reason='measured band '+position
        elif options['graphic']=='band':
            raise ValueError('graphic band cannot fit; revise columns/content allocation or choose none')
        else:
            graphic_reason='content uses available height; retain attached accents only'
    shift=min(h*.12,max(0,(bottom-end)*.30)) if options['vertical_alignment']=='balanced' else 0
    if graphic_rect:
        a,b,c,d=graphic_rect;graphic_rect=(a,b+shift,c,d+shift)
    for role,f,x,y,width,display,index in regions:
        setter.put_fit(role,f,x,y+shift,width,display,index)
    if slide is not None:
        for x,y,width in rules:
            rk.rule(slide,x,y+shift,width,theme.front_ink,lw=.8)
        from pptx.enum.shapes import MSO_CONNECTOR
        from pptx.util import Pt
        from pptx.dml.color import RGBColor
        from pptx.oxml.xmlchemy import OxmlElement
        for diagram in diagrams:
            for i,node in enumerate(diagram['nodes']):
                xx,yy,xx1,yy1=node
                sh=rk.rect(slide,xx,yy+shift,xx1-xx,yy1-yy,fill=None,line=theme.front_ink,lw=.8)
                sh.name=f"information[{diagram['kind']}]/node={i}"
            for a,b in diagram['edges']:
                first,second=diagram['nodes'][a],diagram['nodes'][b]
                cx=(first[0]+first[2])/2
                sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(cx),Pt(first[3]+shift),Pt(cx),Pt(second[1]+shift))
                sh.line.color.rgb=RGBColor.from_string(theme.front_ink.lstrip('#'));sh.line.width=Pt(.8)
                arrow=OxmlElement('a:tailEnd');arrow.set('type','triangle');sh._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln').append(arrow)
                sh.name=f'information[process]/edge={a}:{b}'
    cursor=y1-footer_h+gap if options['footer_anchor']=='edge' else end+shift
    for role,f in foot:
        setter.put_fit(role,f,x0,cursor,w)
        cursor+=f.height+gap
    if image_path and image_rect:
        from pptx.util import Pt
        if not content.get('full_bleed_reason'):
            image_rect=(image_rect[0],image_rect[1]+shift,image_rect[2],image_rect[3]+shift)
        if slide is not None:
            ix0,iy0,ix1,iy1=image_rect
            pic=slide.shapes.add_picture(image_path,Pt(ix0),Pt(iy0),Pt(ix1-ix0),Pt(iy1-iy0))
            pic.name='user-image[duotone]'
            sw,sh=pic.image.size;source_ratio=sw/sh;target_ratio=(ix1-ix0)/(iy1-iy0)
            if content.get('full_bleed_reason'):
                if source_ratio>target_ratio:pic.crop_left=pic.crop_right=(1-target_ratio/source_ratio)/2
                else:pic.crop_top=pic.crop_bottom=(1-source_ratio/target_ratio)/2
            else:
                k=min((ix1-ix0)/sw,(iy1-iy0)/sh)
                pic.width,pic.height=Pt(sw*k),Pt(sh*k)
                pic.left,pic.top=Pt(ix0+(ix1-ix0-sw*k)/2),Pt(iy0+(iy1-iy0-sh*k)/2)
                image_rect=(pic.left.pt,pic.top.pt,pic.left.pt+pic.width.pt,pic.top.pt+pic.height.pt)
            tree=slide.shapes._spTree;tree.remove(pic._element);tree.insert(3,pic._element)
    for diagram in diagrams:
        diagram['nodes']=[[a,b+shift,c,d+shift] for a,b,c,d in diagram['nodes']]
    return {'roles':setter.roles,'image_rect':image_rect,'role_demotions':demotions,'diagrams':diagrams,
            'graphic_rect':graphic_rect,'graphic_reason':graphic_reason,
            'layout':dict(options,columns=columns),'vertical_shift':shift,
            'reading_order':[{'role':r['role'],'block_index':r['block_index'],'order':r['order'],
                              'rect':r['rect'],'shapes':r['shapes']} for r in setter.roles],
            'protagonist':'image' if family=='image_led' else 'key_number' if family=='data_poster' and _present(number_text) else
                           'quote' if family=='manifesto_quote' and quote_text else 'title'}
